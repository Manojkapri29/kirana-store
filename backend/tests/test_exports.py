"""Exports: the reusable engine (CSV/XLSX, formula safety) and the product/inventory/history downloads."""

import csv
import io
from datetime import date, datetime
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app.models.enums import UserRole
from app.services.export_service import (
    CSV_BOM,
    Column,
    ExportFormat,
    Kind,
    neutralize_formula,
    render,
    render_csv,
    render_xlsx,
)
from tests.factories import today_in_shop_timezone

PRODUCTS = "/api/v1/products"
INV = "/api/v1/inventory"
EXPORTS = "/api/v1/exports"

INJECTIONS = [
    "=1+1",
    "+1+1",
    "-2+3",
    "@SUM(A1:A2)",
    '=HYPERLINK("http://evil.example","x")',
    "\t=1+1",
    "\r=1+1",
]


def read_csv(content: bytes) -> list[list[str]]:
    return list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))


def read_xlsx(content: bytes):
    return load_workbook(io.BytesIO(content)).active


def all_text(content: bytes, fmt: str) -> str:
    """Every cell of an export as plain text. XLSX is a compressed zip, so searching raw bytes proves nothing."""
    if fmt == "csv":
        return content.decode("utf-8-sig")
    return "\n".join(
        str(c.value) for row in read_xlsx(content).iter_rows() for c in row if c.value is not None
    )


COLUMNS = [
    Column("name", "Name"),
    Column("price", "Price", Kind.MONEY),
    Column("qty", "Qty", Kind.QUANTITY),
    Column("when", "Date", Kind.DATE),
    Column("at", "Time", Kind.DATETIME),
    Column("n", "Count", Kind.INTEGER),
]
ROW = {
    "name": "Rice",
    "price": Decimal("250.00"),
    "qty": Decimal("2.500"),
    "when": date(2026, 9, 20),
    "at": datetime(2026, 9, 20, 14, 30),
    "n": 7,
}


class TestNeutralizeFormula:
    @pytest.mark.parametrize("text", INJECTIONS)
    def test_dangerous_text_gets_a_leading_apostrophe(self, text):
        assert neutralize_formula(text) == "'" + text

    @pytest.mark.parametrize("text", ["Rice", "5kg Atta", "Amul (1L)", "", "a=b", " =x", "आटा", "₹250"])
    def test_ordinary_text_is_left_alone(self, text):
        assert neutralize_formula(text) == text


class TestCsv:
    def test_utf8_bom_header_and_exact_formatting(self):
        content = render_csv(COLUMNS, [ROW])

        assert content.startswith(CSV_BOM.encode("utf-8"))  # EF BB BF
        assert read_csv(content) == [
            ["Name", "Price", "Qty", "Date", "Time", "Count"],
            ["Rice", "250.00", "2.500", "2026-09-20", "2026-09-20 14:30", "7"],
        ]

    def test_missing_values_are_empty_not_zero(self):
        rows = read_csv(render_csv(COLUMNS, [{"name": "Rice", "price": None}]))

        assert rows[1] == ["Rice", "", "", "", "", ""]

    def test_negative_numbers_are_numbers_and_stay_untouched(self):
        rows = read_csv(
            render_csv(COLUMNS, [{"name": "Adj", "qty": Decimal("-3.000"), "price": Decimal("-0.50")}])
        )

        assert rows[1][1:3] == ["-0.50", "-3.000"]  # not turned into "'-3.000"

    def test_hindi_text_round_trips(self):
        rows = read_csv(render_csv(COLUMNS, [{"name": "बासमती चावल 5 किलो"}]))

        assert rows[1][0] == "बासमती चावल 5 किलो"

    @pytest.mark.parametrize("payload", INJECTIONS)
    def test_formula_injection_is_neutralized(self, payload):
        rows = read_csv(render_csv(COLUMNS, [{"name": payload}]))

        assert rows[1][0] == "'" + payload

    def test_commas_quotes_and_newlines_are_escaped_properly(self):
        tricky = 'Salt, "iodised"\n1kg'

        assert read_csv(render_csv(COLUMNS, [{"name": tricky}]))[1][0] == tricky

    def test_headers_are_neutralized_too(self):
        assert read_csv(render_csv([Column("x", "=cmd")], []))[0] == ["'=cmd"]


class TestXlsx:
    def test_real_excel_types_and_formats(self):
        sheet = read_xlsx(render_xlsx(COLUMNS, [ROW], sheet_name="Test"))
        name, price, qty, when, at, count = sheet[2]

        assert name.value == "Rice"
        assert price.value == 250 and price.number_format == "#,##0.00"
        assert qty.value == Decimal("2.5") and qty.number_format == "#,##0.000"
        assert isinstance(when.value, datetime) and when.value.date() == date(2026, 9, 20)  # a real date cell
        assert when.number_format == "dd/mm/yyyy"
        assert (
            isinstance(at.value, datetime) and at.value.hour == 14 and at.number_format == "dd/mm/yyyy hh:mm"
        )
        assert count.value == 7 and count.data_type == "n"

    def test_header_row_is_bold_and_frozen(self):
        sheet = read_xlsx(render_xlsx(COLUMNS, [ROW], sheet_name="Test"))

        assert [c.value for c in sheet[1]] == ["Name", "Price", "Qty", "Date", "Time", "Count"]
        assert sheet[1][0].font.bold and sheet.freeze_panes == "A2" and sheet.title == "Test"

    @pytest.mark.parametrize("payload", INJECTIONS)
    def test_formula_injection_is_neutralized_and_never_stored_as_a_formula(self, payload):
        cell = read_xlsx(render_xlsx(COLUMNS, [{"name": payload}], sheet_name="T"))["A2"]

        # (XML turns a carriage return into a newline when the file is read back; the prefix is what matters.)
        assert cell.value == ("'" + payload).replace("\r", "\n")
        assert cell.data_type == "s"  # a string, not a formula ("f")

    def test_a_missing_value_is_an_empty_cell(self):
        sheet = read_xlsx(render_xlsx(COLUMNS, [{"name": "Rice", "price": None}], sheet_name="T"))

        assert sheet["B2"].value is None

    def test_negative_numbers_stay_numeric(self):
        sheet = read_xlsx(render_xlsx(COLUMNS, [{"name": "Adj", "qty": Decimal("-3.000")}], sheet_name="T"))

        assert sheet["C2"].value == -3 and sheet["C2"].data_type == "n"

    def test_long_sheet_names_are_shortened_to_excels_limit(self):
        assert len(read_xlsx(render_xlsx(COLUMNS, [], sheet_name="x" * 60)).title) == 31


class TestRender:
    def test_filenames_and_media_types(self):
        csv_file = render(
            ExportFormat.CSV, name="products", columns=COLUMNS, rows=[ROW], on_date=date(2026, 9, 20)
        )
        xlsx_file = render(
            ExportFormat.XLSX,
            name="inventory_history",
            columns=COLUMNS,
            rows=[ROW],
            on_date=date(2026, 9, 20),
        )

        assert (csv_file.filename, csv_file.media_type) == (
            "products_2026-09-20.csv",
            "text/csv; charset=utf-8",
        )
        assert xlsx_file.filename == "inventory_history_2026-09-20.xlsx"
        assert xlsx_file.media_type.endswith("spreadsheetml.sheet")
        assert read_xlsx(xlsx_file.content).title == "Inventory History"


# --- The real downloads ----------------------------------------------------------------------------


def product(client, tenant, units, **overrides):
    data = {
        "sku": "RICE", "name": "Rice 5kg", "category_id": tenant.category.id,
        "unit_id": units["pcs"], "selling_price": "250", "reorder_level": "5",
    }  # fmt: skip
    data.update(overrides)
    response = client.post(PRODUCTS, json=data)
    assert response.status_code == 201, response.text
    return response.json()["product"]


@pytest.fixture
def shops(client_a, client_b, tenant_a, tenant_b, units):
    """Shop A has two products (one with stock and an unsafe name); Shop B has one."""
    mine = product(
        client_a,
        tenant_a,
        units,
        sku="A-RICE",
        name="Rice",
        brand="India Gate",
        mrp="300",
        opening_stock="20",
        opening_stock_cost="200",
    )
    product(
        client_a,
        tenant_a,
        units,
        sku="A-EVIL",
        name='=HYPERLINK("http://evil.example","click")',
        barcode="8901234567890",
    )
    product(client_b, tenant_b, units, sku="B-SECRET", name="Other shop's product", opening_stock="99")
    return mine


class TestProductExport:
    def test_csv(self, client_a, shops):
        response = client_a.get(f"{EXPORTS}/products", params={"format": "csv"})

        assert response.status_code == 200
        assert response.headers["content-type"] == "text/csv; charset=utf-8"
        assert (
            response.headers["content-disposition"]
            == f'attachment; filename="products_{today_in_shop_timezone().isoformat()}.csv"'
        )
        assert response.headers["cache-control"] == "no-store"
        assert response.content.startswith(b"\xef\xbb\xbf")
        rows = read_csv(response.content)
        header = rows[0]
        assert header[:3] == ["SKU", "Product", "Brand"]
        by_sku = {r[0]: dict(zip(header, r, strict=True)) for r in rows[1:]}
        assert set(by_sku) == {"A-RICE", "A-EVIL"}
        rice = by_sku["A-RICE"]
        assert (rice["MRP"], rice["Selling Price"], rice["Average Cost"], rice["Current Stock"]) == (
            "300.00",
            "250.00",
            "200.00",
            "20.000",
        )
        assert (rice["Stock Status"], rice["Status"], rice["Purchase Price"]) == ("In Stock", "Active", "")

    def test_xlsx(self, client_a, shops):
        response = client_a.get(f"{EXPORTS}/products", params={"format": "xlsx"})

        assert response.status_code == 200
        assert response.headers["content-type"].endswith("spreadsheetml.sheet")
        sheet = read_xlsx(response.content)
        header = [c.value for c in sheet[1]]
        rows = {
            r[0].value: dict(zip(header, (c.value for c in r), strict=True))
            for r in sheet.iter_rows(min_row=2)
        }
        assert set(rows) == {"A-RICE", "A-EVIL"}
        assert rows["A-RICE"]["Selling Price"] == 250 and rows["A-RICE"]["Current Stock"] == 20
        created = next(c for c in sheet[2] if isinstance(c.value, datetime))
        assert created.number_format == "dd/mm/yyyy hh:mm"  # a real Excel date-time cell

    @pytest.mark.parametrize("fmt", ["csv", "xlsx"])
    def test_only_the_callers_shop_is_exported(self, client_a, client_b, shops, fmt):
        mine = all_text(client_a.get(f"{EXPORTS}/products", params={"format": fmt}).content, fmt)
        theirs = all_text(client_b.get(f"{EXPORTS}/products", params={"format": fmt}).content, fmt)

        assert "A-RICE" in mine and "B-SECRET" not in mine and "Other shop" not in mine
        assert "B-SECRET" in theirs and "A-RICE" not in theirs and "A-EVIL" not in theirs

    def test_formula_injection_in_a_product_name_is_neutralized_in_both_formats(self, client_a, shops):
        evil = '=HYPERLINK("http://evil.example","click")'

        csv_rows = read_csv(client_a.get(f"{EXPORTS}/products", params={"format": "csv"}).content)
        xlsx_sheet = read_xlsx(client_a.get(f"{EXPORTS}/products", params={"format": "xlsx"}).content)

        assert "'" + evil in [r[1] for r in csv_rows[1:]]
        assert all(r[1] != evil for r in csv_rows)
        cell = next(
            c for row in xlsx_sheet.iter_rows(min_row=2) for c in row if c.value and evil in str(c.value)
        )
        assert cell.value == "'" + evil and cell.data_type == "s"

    def test_status_filter_and_default(self, client_a, shops):
        client_a.post(f"{PRODUCTS}/{shops['id']}/deactivate")

        active = read_csv(client_a.get(f"{EXPORTS}/products").content)
        inactive = read_csv(client_a.get(f"{EXPORTS}/products", params={"status": "inactive"}).content)
        everything = read_csv(client_a.get(f"{EXPORTS}/products", params={"status": "all"}).content)

        assert [r[0] for r in active[1:]] == ["A-EVIL"]
        assert [r[0] for r in inactive[1:]] == ["A-RICE"]
        assert len(everything) == 3

    def test_default_format_is_csv_and_a_bad_format_is_refused(self, client_a, shops):
        assert client_a.get(f"{EXPORTS}/products").headers["content-type"].startswith("text/csv")
        assert client_a.get(f"{EXPORTS}/products", params={"format": "pdf"}).status_code == 422


class TestInventoryExport:
    def test_csv_matches_the_inventory_screen(self, client_a, shops):
        rows = read_csv(client_a.get(f"{EXPORTS}/inventory").content)
        header = rows[0]
        by_sku = {r[0]: dict(zip(header, r, strict=True)) for r in rows[1:]}

        assert set(by_sku) == {"A-RICE", "A-EVIL"}
        assert (
            by_sku["A-RICE"]["Current Stock"] == "20.000" and by_sku["A-RICE"]["Stock Status"] == "In Stock"
        )
        assert (
            by_sku["A-EVIL"]["Current Stock"] == "0.000"
            and by_sku["A-EVIL"]["Stock Status"] == "Out of Stock"
        )

    def test_stock_status_filter(self, client_a, shops):
        rows = read_csv(client_a.get(f"{EXPORTS}/inventory", params={"stock_status": "OUT_OF_STOCK"}).content)

        assert [r[0] for r in rows[1:]] == ["A-EVIL"]

    def test_xlsx_is_shop_scoped(self, client_a, shops):
        content = client_a.get(f"{EXPORTS}/inventory", params={"format": "xlsx"}).content

        assert "A-RICE" in all_text(content, "xlsx") and "B-SECRET" not in all_text(content, "xlsx")
        assert read_xlsx(content).max_row == 3  # header + Shop A's two products


class TestInventoryHistoryExport:
    def test_csv_has_the_ledger_with_running_balance(self, client_a, shops):
        rows = read_csv(client_a.get(f"{EXPORTS}/inventory-history").content)
        header = rows[0]
        row = dict(zip(header, rows[1], strict=True))

        assert len(rows) == 2  # only Shop A's single OPENING row
        assert (row["Type"], row["Quantity Change"], row["Balance After"], row["Unit Cost"]) == (
            "OPENING",
            "20.000",
            "20.000",
            "200.00",
        )
        assert (row["SKU"], row["Source"], row["Recorded By"]) == ("A-RICE", "PRODUCT", "Test Owner")
        assert row["Date"] == today_in_shop_timezone().isoformat()

    def test_xlsx_dates_are_real_date_cells(self, client_a, shops):
        sheet = read_xlsx(client_a.get(f"{EXPORTS}/inventory-history", params={"format": "xlsx"}).content)
        first = sheet[2][0]

        assert isinstance(first.value, datetime) and first.number_format == "dd/mm/yyyy"
        assert sheet["E2"].value == 20 and sheet["E2"].number_format == "#,##0.000"

    def test_is_shop_scoped_and_filterable(self, client_a, client_b, shops):
        mine = all_text(client_a.get(f"{EXPORTS}/inventory-history").content, "csv")
        theirs = all_text(client_b.get(f"{EXPORTS}/inventory-history").content, "csv")

        assert "A-RICE" in mine and "B-SECRET" not in mine
        assert "B-SECRET" in theirs and "A-RICE" not in theirs
        assert (
            len(read_csv(client_a.get(f"{EXPORTS}/inventory-history", params={"product_id": 999999}).content))
            == 1
        )

    def test_the_other_shops_product_id_exports_nothing(self, client_a, client_b, tenant_b, units, shops):
        theirs = product(client_b, tenant_b, units, sku="B-2")

        rows = read_csv(
            client_a.get(f"{EXPORTS}/inventory-history", params={"product_id": theirs["id"]}).content
        )

        assert len(rows) == 1  # header only


class TestExportAccess:
    def test_only_the_owner_can_export(self, make_client, tenant_a):
        staff = make_client(tenant_a, role=UserRole.STAFF)

        for path in ("products", "inventory", "inventory-history"):
            assert staff.get(f"{EXPORTS}/{path}").status_code == 403

    def test_staff_can_still_use_the_normal_screens(self, make_client, tenant_a):
        assert make_client(tenant_a, role=UserRole.STAFF).get(PRODUCTS).status_code == 200
