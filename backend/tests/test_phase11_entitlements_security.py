"""Plan entitlements and usage for the Phase 11 features, exports, and payload/IDOR/search security checks."""

import csv
import io
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app.models import PlanFeature
from app.services import entitlement_service
from tests.test_promotions import shop as shop  # noqa: F401
from tests.test_purchases_api import make_product, make_supplier
from tests.test_sales_api import item, make_customer, sold  # noqa: F401

EXPORTS = "/api/v1/exports"


def set_feature(session_factory, plan_code, key, *, enabled=None, limit=None):
    with session_factory() as s, s.begin():
        plan = entitlement_service.get_plan(s, plan_code)
        values = {}
        if enabled is not None:
            values["enabled"] = enabled
        if limit is not None:
            values["limit_value"] = limit
        s.execute(
            PlanFeature.__table__.update()
            .where(PlanFeature.plan_id == plan.id, PlanFeature.feature_key == key)
            .values(**values)
        )


class TestEntitlementsAndUsage:
    def test_usage_lists_every_metered_thing_with_used_limit_remaining(self, client_a):
        body = client_a.get("/api/v1/account/usage").json()
        assert body["plan_code"] and body["period"]
        assert {
            "max_products",
            "max_ai_requests_per_month",
            "max_exports_per_month",
            "max_image_analyses_per_month",
        } <= set(body["limits"])
        for item_ in body["limits"].values():
            assert set(item_) >= {"used", "limit", "remaining"}

    def test_each_export_counts_and_the_plan_can_cap_them(
        self, client_a, tenant_a, give_plan, session_factory
    ):
        give_plan(tenant_a, "free")
        set_feature(session_factory, "free", "max_exports_per_month", limit=2)
        assert client_a.get(f"{EXPORTS}/products").status_code == 200
        assert client_a.get(f"{EXPORTS}/customers").status_code == 200
        blocked = client_a.get(f"{EXPORTS}/products")
        assert blocked.status_code == 403 and blocked.json()["category"] == "plan_limit"
        assert "exports this month" in blocked.json()["message"]
        usage = client_a.get("/api/v1/account/usage").json()["limits"]["max_exports_per_month"]
        assert (usage["used"], usage["limit"], usage["remaining"]) == (2, 2, 0)
        assert client_a.get("/api/v1/customers").status_code == 200  # nothing else is affected

    def test_usage_is_counted_per_shop(
        self, client_a, client_b, tenant_a, tenant_b, give_plan, session_factory
    ):
        give_plan(tenant_a, "free")
        give_plan(tenant_b, "free")
        set_feature(session_factory, "free", "max_exports_per_month", limit=1)
        assert client_a.get(f"{EXPORTS}/products").status_code == 200
        assert client_a.get(f"{EXPORTS}/products").status_code == 403
        assert client_b.get(f"{EXPORTS}/products").status_code == 200  # shop B's allowance is its own

    def test_a_plan_without_the_export_feature_refuses_and_names_it(
        self, client_a, tenant_a, give_plan, session_factory
    ):
        give_plan(tenant_a, "free")
        set_feature(session_factory, "free", "exports", enabled=False)
        response = client_a.get(f"{EXPORTS}/products")
        assert response.status_code == 403 and "data exports" in response.json()["message"]

    def test_the_owner_is_told_once_when_an_allowance_is_used_up(
        self, client_a, tenant_a, give_plan, session_factory
    ):
        give_plan(tenant_a, "free")
        set_feature(session_factory, "free", "max_exports_per_month", limit=1)
        client_a.get(f"{EXPORTS}/products")
        rows = client_a.get("/api/v1/notifications").json()["items"]
        assert [r["event_type"] for r in rows] == ["SUBSCRIPTION_LIMIT"]

    def test_photo_analyses_are_counted_and_capped(self, client_a, tenant_a, give_plan, session_factory):
        give_plan(tenant_a, "free")
        set_feature(session_factory, "free", "max_image_analyses_per_month", limit=0)
        response = client_a.post(
            "/api/v1/image-intelligence/analyze", files={"image": ("a.png", b"x", "image/png")}
        )
        assert response.status_code == 403 and response.json()["category"] == "plan_limit"

    def test_a_product_limit_is_a_clear_refusal_and_deletes_nothing(
        self, client_a, tenant_a, units, give_plan, session_factory
    ):
        give_plan(tenant_a, "free")
        set_feature(session_factory, "free", "max_products", limit=1)
        make_product(client_a, tenant_a, units, "ONE")
        blocked = client_a.post(
            "/api/v1/products",
            json={
                "name": "Two",
                "sku": "TWO",
                "unit_id": units["pcs"],
                "category_id": tenant_a.category.id,
                "selling_price": "5",
            },
        )
        assert blocked.status_code == 403 and blocked.json()["category"] == "plan_limit"
        assert (
            len(client_a.get("/api/v1/products").json()["items"]) == 1
        )  # what exists stays exactly as it was


class TestNoInternalDetailInExports:
    def test_the_suppliers_export_is_shop_scoped_utf8_bom_and_formula_safe(self, client_a, client_b):
        make_supplier(client_a, '=HYPERLINK("http://evil.example","x")')
        make_supplier(client_b, "Other Shop Supplier")
        response = client_a.get(f"{EXPORTS}/suppliers")
        assert response.status_code == 200 and response.content.startswith(b"\xef\xbb\xbf")
        assert response.headers["content-disposition"].startswith('attachment; filename="suppliers_')
        rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
        names = [r[0] for r in rows[1:]]
        assert (
            names == ['\'=HYPERLINK("http://evil.example","x")']
            and "Other Shop Supplier" not in response.text
        )
        assert rows[0][0] == "Supplier" and "id" not in [h.lower() for h in rows[0]]

    def test_the_suppliers_export_also_comes_as_a_workbook(self, client_a):
        make_supplier(client_a, "Sharma")
        sheet = load_workbook(
            io.BytesIO(client_a.get(f"{EXPORTS}/suppliers", params={"format": "xlsx"}).content)
        ).active
        created = sheet.cell(row=2, column=10).value
        assert sheet.cell(row=2, column=1).value == "Sharma" and hasattr(
            created, "year"
        )  # a real date cell, not text

    def test_exports_are_owner_only_and_never_cached(self, client_a, tenant_a, make_client):
        from app.models.enums import UserRole

        staff = make_client(tenant_a, UserRole.STAFF)
        assert staff.get(f"{EXPORTS}/suppliers").status_code == 403
        assert client_a.get(f"{EXPORTS}/suppliers").headers["cache-control"] == "no-store"


class TestPayloadManipulation:
    def test_a_shop_id_in_a_payload_is_never_trusted(self, client_a, client_b, tenant_a, tenant_b):
        made = client_a.post("/api/v1/customers", json={"name": "Sneaky", "shop_id": tenant_b.shop.id})
        assert made.status_code in (201, 422)
        assert client_b.get("/api/v1/customers", params={"q": "Sneaky"}).json()["total"] == 0
        if made.status_code == 201:
            assert client_a.get("/api/v1/customers", params={"q": "Sneaky"}).json()["total"] == 1

    def test_stock_effects_and_prices_cannot_be_forced_through_a_sale_payload(self, client_a, shop):
        forged = client_a.post(
            "/api/v1/sales",
            json={"items": [item(shop["rice"], "2", unit_cost="1", cogs_amount="0", line_total="1")]},
        )
        assert forged.status_code == 422  # unknown fields are refused, not silently ignored
        s = sold(client_a, [item(shop["rice"], "2")])
        assert (
            client_a.get(f"/api/v1/sales/{s['id']}").json()["total_amount"] == "100.00"
        )  # the price comes from the product
        assert (
            Decimal(client_a.get(f"/api/v1/inventory/products/{shop['rice']['id']}").json()["current_stock"])
            == 18
        )

    def test_a_customer_balance_cannot_be_written_directly(self, client_a):
        made = client_a.post(
            "/api/v1/customers", json={"name": "Ram", "balance": "-9999", "outstanding": "0"}
        )
        assert made.status_code in (201, 422)
        if made.status_code == 201:
            customer = made.json()["customer"]
            assert client_a.get(f"/api/v1/customers/{customer['id']}/balance").json()["balance"] == "0.00"

    def test_another_shops_records_are_a_404_not_a_403(self, client_a, client_b, shop):
        product_id = shop["rice"]["id"]
        assert client_b.get(f"/api/v1/products/{product_id}").status_code == 404
        assert client_b.get(f"/api/v1/inventory/products/{product_id}").status_code == 404
        assert client_b.get(f"{EXPORTS}/customers/999999/ledger").status_code == 404


class TestSearchSafety:
    @pytest.mark.parametrize("q", ["%", "_", "%%%%", "a%b", "' OR 1=1 --", "\\", "%_%"])
    def test_wildcards_and_sql_are_plain_text(self, client_a, tenant_a, units, q):
        make_product(client_a, tenant_a, units, "PLAIN")
        for path in ("/api/v1/products", "/api/v1/customers", "/api/v1/suppliers"):
            response = client_a.get(path, params={"q": q})
            assert response.status_code == 200, (path, q)
            assert response.json()["total"] == 0  # "%" does not mean "everything"

    @pytest.mark.parametrize(
        "path", ["/api/v1/products", "/api/v1/customers", "/api/v1/suppliers", f"{EXPORTS}/products"]
    )
    def test_an_oversized_query_is_refused_not_run(self, client_a, path):
        assert client_a.get(path, params={"q": "x" * 5000}).status_code == 422

    def test_search_never_crosses_shops(self, client_a, client_b, tenant_a, units):
        make_product(client_a, tenant_a, units, "SECRET1", name="Secret Item")
        assert client_b.get("/api/v1/products", params={"q": "Secret"}).json()["total"] == 0


class TestPaginationConvention:
    @pytest.mark.parametrize(
        "path",
        [
            "/api/v1/products",
            "/api/v1/customers",
            "/api/v1/suppliers",
            "/api/v1/sales",
            "/api/v1/purchases",
            "/api/v1/notifications",
            "/api/v1/audit-log",
        ],
    )
    def test_lists_report_total_and_bounded_pages(self, client_a, path):
        response = client_a.get(path, params={"limit": 5, "offset": 0})
        assert response.status_code == 200
        body = response.json()
        assert "total" in body and isinstance(body["items"], list) and len(body["items"]) <= 5

    @pytest.mark.parametrize("path", ["/api/v1/products", "/api/v1/customers"])
    def test_a_huge_page_size_is_refused(self, client_a, path):
        assert client_a.get(path, params={"limit": 100000}).status_code == 422


class TestOpenApiHasNoSecrets:
    def test_docs_describe_the_api_and_hold_no_configuration_values(self, client_a):
        spec = client_a.get("/openapi.json").text
        for word in ("secret_key", "api_key=", "sk-", "password", "KIRANA_"):
            assert word.lower() not in spec.lower().replace("secret_key_hint", ""), word


def test_deletes_do_not_exist_for_financial_documents(client_a, shop):
    s = sold(client_a, [item(shop["rice"], "1")])
    for path in (
        f"/api/v1/sales/{s['id']}",
        "/api/v1/customers/1",
        "/api/v1/products/1",
        "/api/v1/admin/backups/x",
    ):
        assert client_a.delete(path).status_code in (404, 405)
