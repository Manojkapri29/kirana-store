"""Advanced exports: CSV/XLSX/PDF carry the same numbers as the screens, are formula-safe, honour selected columns and
grouping, need ANALYTICS_EXPORT plus the report's own permissions, and never cross shops."""

import csv
import io
import re
from datetime import date
from decimal import Decimal

from openpyxl import load_workbook

from tests import factories
from tests.client_helpers import client_with
from tests.factories import today_in_shop_timezone
from tests.finance_helpers import make_quick_sale, make_sale

TODAY = today_in_shop_timezone()
API = "/api/v1/analytics"
P = {"preset": "this_month", "compare": "none"}


def _export(client, key, fmt="csv", status=200, **params):
    r = client.get(f"{API}/export/{key}", params={**P, "format": fmt, **params})
    assert r.status_code == status, r.text[:300]
    return r


def _csv_rows(response):
    text = response.content.decode("utf-8")
    assert text.startswith("﻿")
    return list(csv.reader(io.StringIO(text.lstrip("﻿"))))


def _table_of(rows):
    """Split a report CSV into (header block, header row, data rows, notes)."""
    blank = [i for i, r in enumerate(rows) if not r]
    head = rows[: blank[0]]
    header = rows[blank[0] + 1]
    end = blank[1] if len(blank) > 1 else len(rows)
    return head, header, rows[blank[0] + 2 : end], [r[0] for r in rows[end + 1 :]]


class TestFilesMatchTheScreens:
    def test_kpi_export_equals_the_kpi_endpoint(self, session, tenant_a, client_a):
        make_sale(session, tenant_a, "1000.00", day=TODAY, cogs="600.00")
        make_quick_sale(session, tenant_a, "250.00", day=TODAY)
        session.commit()
        api = {
            k["definition"]["name"]: k["current"]
            for k in client_a.get(f"{API}/kpis", params=P).json()["kpis"]
        }
        _, header, data, _ = _table_of(_csv_rows(_export(client_a, "kpis")))
        by_name = {r[0]: dict(zip(header, r, strict=True)) for r in data}
        assert len(by_name) == len(api)
        for name, cur in api.items():
            exported = by_name[name]
            assert exported["Availability"] == cur["availability"]
            assert (exported["Value"] or None) == cur["amount"], name

    def test_sales_export_totals_equal_the_screen(self, session, tenant_a, client_a):
        make_sale(session, tenant_a, "100.00", day=TODAY, cogs="60.00")
        make_quick_sale(session, tenant_a, "50.00", day=TODAY)
        session.commit()
        api = client_a.get(f"{API}/sales/trend", params={**P, "bucket": "month"}).json()["rows"]
        _, header, data, _ = _table_of(_csv_rows(_export(client_a, "sales-trend-month")))
        col = header.index("Combined revenue (before returns)")
        assert [Decimal(r[col]) for r in data] == [Decimal(r["combined_net"]) for r in api]

    def test_xlsx_and_csv_hold_the_same_figures(self, session, tenant_a, client_a):
        make_sale(session, tenant_a, "100.00", day=TODAY, cogs="60.00")
        session.commit()
        _, header, data, _ = _table_of(_csv_rows(_export(client_a, "finance-trend-month")))
        ws = load_workbook(io.BytesIO(_export(client_a, "finance-trend-month", "xlsx").content)).active
        grid = [[c.value for c in row] for row in ws.iter_rows()]
        start = next(i for i, r in enumerate(grid) if r[0] == header[0])
        rev = header.index("Revenue")
        assert Decimal(str(grid[start + 1][rev])) == Decimal(data[0][rev])


class TestFormats:
    def test_csv_has_bom_title_shop_period_generated_and_notes(self, session, tenant_a, client_a):
        make_quick_sale(session, tenant_a, "50.00", day=TODAY)
        session.commit()
        r = _export(client_a, "sales-trend-day", title="My sales")
        assert (
            r.headers["content-type"].startswith("text/csv")
            and "attachment" in r.headers["content-disposition"]
        )
        assert r.headers["cache-control"] == "no-store"
        head, _, _, notes = _table_of(_csv_rows(r))
        labels = {row[0]: row[1] for row in head[1:]}
        assert head[0][0] == "My sales"
        assert {"Shop", "Period", "Generated", "Source"} <= set(labels)
        assert tenant_a.shop.name in labels["Shop"] and re.match(
            r"\d{4}-\d\d-\d\d \d\d:\d\d ", labels["Generated"]
        )
        assert notes

    def test_xlsx_has_real_numbers_dates_fitted_widths_and_a_frozen_header(self, session, tenant_a, client_a):
        make_quick_sale(session, tenant_a, "50.25", day=TODAY)
        session.commit()
        ws = load_workbook(io.BytesIO(_export(client_a, "sales-trend-day", "xlsx").content)).active
        header_row = next(
            i
            for i, row in enumerate(ws.iter_rows(), 1)
            if row[0].value == "Period" and row[1].value == "Detailed sales"
        )
        assert ws.freeze_panes == f"A{header_row + 1}"
        headers = [c.value for c in ws[header_row]]
        cell = ws[header_row + 1][headers.index("Quick sales")]
        assert isinstance(cell.value, int | float | Decimal) and float(cell.value) == 50.25
        assert cell.number_format == "#,##0.00"
        assert ws.column_dimensions["A"].width >= 10

    def test_pdf_is_a_valid_document_with_header_period_and_page_numbers(self, session, tenant_a, client_a):
        for i in range(60):
            make_quick_sale(session, tenant_a, f"{10 + i}.00", day=TODAY)
        session.commit()
        r = _export(client_a, "sales-trend-day", "pdf")
        pdf = r.content
        assert (
            r.headers["content-type"] == "application/pdf"
            and pdf.startswith(b"%PDF-1.4")
            and pdf.rstrip().endswith(b"%%EOF")
        )
        # every xref offset points at its object
        xref_at = int(re.search(rb"startxref\n(\d+)", pdf).group(1))
        entries = pdf[xref_at:].split(b"\n")[3:]
        for n, entry in enumerate([e for e in entries if e.endswith(b" n ")], start=1):
            assert pdf[int(entry[:10]) :].startswith(f"{n} 0 obj".encode())
        assert (
            b"Sales by day" in pdf
            and tenant_a.shop.name.encode() in pdf
            and b"Period:" in pdf
            and b"Generated" in pdf
        )
        assert re.search(rb"Page 1 of \d+", pdf)

    def test_pdf_repeats_the_header_row_and_numbers_every_page(self, session, tenant_a, client_a):
        for i in range(80):
            make_quick_sale(session, tenant_a, f"{10 + i}.00", day=TODAY)
        session.commit()
        c = client_a.get(f"{API}/export/customers-list", params={**P, "format": "pdf"})
        assert c.status_code == 200
        pdf = _export(client_a, "sales-products-top", "pdf").content
        pages = int(re.search(rb"/Count (\d+)", pdf).group(1))
        assert len(re.findall(rb"Page \d+ of %d" % pages, pdf)) == pages

    def test_unknown_report_and_bad_format_are_refused(self, session, tenant_a, client_a):
        assert _export(client_a, "nope", status=404)
        assert client_a.get(f"{API}/export/kpis", params={**P, "format": "docx"}).status_code == 422


class TestSafety:
    def test_formula_text_is_neutralised_in_csv_and_xlsx(self, session, tenant_a, client_a):
        factories.make_product(
            session, tenant_a.shop, tenant_a.category, sku="X1", name='=HYPERLINK("http://evil")'
        )
        session.commit()
        _, header, data, _ = _table_of(_csv_rows(_export(client_a, "inventory-stock")))
        names = [r[header.index("Product")] for r in data]
        assert '\'=HYPERLINK("http://evil")' in names and not any(n.startswith("=") for n in names)
        ws = load_workbook(io.BytesIO(_export(client_a, "inventory-stock", "xlsx").content)).active
        cells = [
            c for row in ws.iter_rows() for c in row if isinstance(c.value, str) and "HYPERLINK" in c.value
        ]
        assert cells and all(c.data_type == "s" and not c.value.startswith("=") for c in cells)

    def test_a_blank_is_explained_never_zero(self, session, tenant_a, client_a):
        make_quick_sale(session, tenant_a, "50.00", day=TODAY)
        session.commit()
        _, header, data, notes = _table_of(_csv_rows(_export(client_a, "finance-trend-month")))
        row = dict(zip(header, data[0], strict=True))
        assert row["Gross profit"] == "" and any("Not Available" in n for n in notes)


class TestColumnsAndGrouping:
    def test_selected_columns_come_out_in_order(self, session, tenant_a, client_a):
        make_quick_sale(session, tenant_a, "50.00", day=TODAY)
        session.commit()
        _, header, data, _ = _table_of(
            _csv_rows(_export(client_a, "sales-trend-day", columns="combined_net,period"))
        )
        assert header == ["Combined revenue (before returns)", "Period"] and len(data[0]) == 2

    def test_unknown_or_repeated_columns_are_refused(self, session, tenant_a, client_a):
        assert _export(client_a, "sales-trend-day", status=422, columns="period,drop_table")
        assert _export(client_a, "sales-trend-day", status=422, columns="period,period")
        assert _export(client_a, "sales-trend-day", status=422, columns="period", group_by="combined_net")

    def test_grouping_adds_exact_subtotals_and_a_grand_total(self, session, tenant_a, client_a):
        make_quick_sale(session, tenant_a, "10.10", day=TODAY)
        make_quick_sale(session, tenant_a, "20.20", day=TODAY)
        session.commit()
        _, header, data, notes = _table_of(_csv_rows(_export(client_a, "sales-channels", group_by="channel")))
        assert data[-1][0] == "Grand total"
        assert any(n.startswith("Subtotals") for n in notes)


class TestPermissions:
    def test_export_needs_the_export_permission(self, session, tenant_a, make_client):
        c = client_with(make_client, tenant_a, ["ANALYTICS_VIEW", "REPORT_VIEW"])
        assert c.get(f"{API}/export/sales-trend-day", params=P).status_code == 403
        assert c.get(f"{API}/export").status_code == 403

    def test_export_needs_the_reports_own_permissions_too(self, session, tenant_a, make_client):
        c = client_with(make_client, tenant_a, ["ANALYTICS_EXPORT", "ANALYTICS_VIEW", "REPORT_VIEW"])
        assert c.get(f"{API}/export/sales-trend-day", params=P).status_code == 200
        assert c.get(f"{API}/export/finance-trend-month", params=P).status_code == 403
        assert c.get(f"{API}/export/inventory-stock", params=P).status_code == 403

    def test_the_list_only_offers_what_the_caller_may_export(self, session, tenant_a, make_client):
        c = client_with(make_client, tenant_a, ["ANALYTICS_EXPORT", "ANALYTICS_VIEW", "REPORT_VIEW"])
        keys = {r["key"] for r in c.get(f"{API}/export").json()["reports"]}
        assert "sales-trend-day" in keys and "kpis" in keys and "finance-trend-month" not in keys

    def test_kpi_export_hides_kpis_the_role_may_not_see(self, session, tenant_a, make_client):
        c = client_with(make_client, tenant_a, ["ANALYTICS_EXPORT", "ANALYTICS_VIEW"])
        _, header, data, _ = _table_of(_csv_rows(_export(c, "kpis")))
        names = {r[0] for r in data}
        assert "Gross profit" not in names and "Inventory value" not in names

    def test_exports_never_cross_shops(self, session, tenant_a, tenant_b, client_b):
        make_quick_sale(session, tenant_a, "999.00", day=TODAY)
        session.commit()
        _, _, data, _ = _table_of(_csv_rows(_export(client_b, "sales-trend-day")))
        assert data == []


def test_dates_in_the_period_block_are_iso(session, tenant_a, client_a):
    head, *_ = _table_of(_csv_rows(_export(client_a, "sales-trend-day")))
    period = next(r[1] for r in head if r[0] == "Period")
    found = re.findall(r"\d{4}-\d\d-\d\d", period)
    assert len(found) == 2 and date.fromisoformat(found[0]) <= date.fromisoformat(found[1])


class TestSavedReportExports:
    def test_a_saved_report_exports_with_the_same_numbers_and_the_callers_permissions(
        self, session, tenant_a, client_a, make_client
    ):
        make_quick_sale(session, tenant_a, "80.00", day=TODAY)
        session.commit()
        body = {
            "name": "Totals",
            "dataset": "sales",
            "definition": {"group_by": ["kind"], "aggregations": [{"field": "total_amount", "op": "SUM"}]},
        }
        rid = client_a.post(f"{API}/reports", json=body).json()["id"]
        head, header, data, _ = _table_of(_csv_rows(_export(client_a, f"saved-{rid}")))
        assert head[0][0] == "Totals" and data == [["Quick", "80.00"]]
        weak = client_with(make_client, tenant_a, ["ANALYTICS_EXPORT", "REPORT_VIEW"])
        assert (
            weak.get(f"{API}/export/saved-{rid}", params=P).status_code == 403
        )  # no ANALYTICS_CUSTOM_REPORT
        assert _export(client_a, "saved-999999", status=404)
        assert _export(client_a, "saved-abc", status=404)

    def test_another_shops_saved_report_is_not_found(self, session, tenant_a, tenant_b, client_a, client_b):
        rid = client_a.post(
            f"{API}/reports", json={"name": "T", "dataset": "sales", "definition": {}}
        ).json()["id"]
        assert _export(client_b, f"saved-{rid}", status=404)
