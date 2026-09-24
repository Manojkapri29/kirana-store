"""Reusable CSV / XLSX export engine (BUSINESS_RULES X1-X4).

A dataset is described by a list of `Column`s and a list of row dicts. This module knows nothing about
products or stock, so purchases, sales, returns, expenses, khata and reports can reuse it in later phases:
each just supplies its columns and rows.

Safety and Excel-friendliness:
  * CSV is UTF-8 **with BOM**, so Excel shows Hindi text and the rupee sign correctly.
  * XLSX cells are real Excel types: numbers are numbers, dates are date cells (formatted dd/mm/yyyy).
  * **Formula injection**: a text value that starts with `=`, `+`, `-`, `@`, tab or carriage return could be
    run as a formula when the file is opened in Excel/Sheets. Such text gets a leading apostrophe, which
    makes spreadsheets treat it as plain text. Only *text* is changed; genuine numbers such as -5 are not.
  * Amounts are written from `Decimal`, never through a float.
Callers must pass only rows that already belong to the current shop; this module does not query anything.
"""

import csv
import io
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from enum import StrEnum
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

CSV_BOM = "﻿"
DANGEROUS_PREFIXES = ("=", "+", "-", "@", "\t", "\r")

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_MEDIA_TYPE = "text/csv; charset=utf-8"
PDF_MEDIA_TYPE = "application/pdf"


class ExportFormat(StrEnum):
    CSV = "csv"
    XLSX = "xlsx"


class ReportFormat(StrEnum):
    """The formats of an analytics report file. The older `/exports/*` downloads stay CSV and XLSX only (`ExportFormat`)."""

    CSV = "csv"
    XLSX = "xlsx"
    PDF = "pdf"


class Kind(StrEnum):
    """How a column is stored in the file."""

    TEXT = "text"
    INTEGER = "integer"
    MONEY = "money"  # 2 decimals
    QUANTITY = "quantity"  # 3 decimals
    PERCENT = "percent"  # 2 decimals, a percentage or a ratio
    DECIMAL = "decimal"  # a number whose unit varies by row (a KPI table): plain, no thousands separator
    DATE = "date"
    DATETIME = "datetime"


@dataclass(frozen=True)
class Column:
    key: str  # key in each row dict
    header: str
    kind: Kind = Kind.TEXT


@dataclass(frozen=True)
class ExportFile:
    filename: str
    media_type: str
    content: bytes


_NUMBER_FORMATS = {
    Kind.INTEGER: "0",
    Kind.MONEY: "#,##0.00",
    Kind.QUANTITY: "#,##0.000",
    Kind.PERCENT: "0.00",
    Kind.DECIMAL: "General",
    Kind.DATE: "dd/mm/yyyy",
    Kind.DATETIME: "dd/mm/yyyy hh:mm",
}


def neutralize_formula(value: str) -> str:
    """Make text safe to open in a spreadsheet. Text that could be read as a formula gets a leading '."""
    return "'" + value if value.startswith(DANGEROUS_PREFIXES) else value


def _clean(value: Any) -> Any:
    return neutralize_formula(value) if isinstance(value, str) else value


def _csv_text(value: Any, kind: Kind) -> str:
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        return neutralize_formula(value)
    return str(value)


def render_csv(columns: Sequence[Column], rows: Iterable[Mapping[str, Any]]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([neutralize_formula(column.header) for column in columns])
    for row in rows:
        writer.writerow([_csv_text(row.get(column.key), column.kind) for column in columns])
    return (CSV_BOM + buffer.getvalue()).encode("utf-8")


def render_xlsx(columns: Sequence[Column], rows: Iterable[Mapping[str, Any]], *, sheet_name: str) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name[:31]  # Excel's limit

    sheet.append([neutralize_formula(column.header) for column in columns])
    header_fill = PatternFill("solid", start_color="E2F0E8")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"

    widths = [len(column.header) for column in columns]
    for row in rows:
        values = [_clean(row.get(column.key)) for column in columns]
        sheet.append(values)
        for index, (column, value) in enumerate(zip(columns, values, strict=True)):
            cell = sheet.cell(row=sheet.max_row, column=index + 1)
            if value is None:
                continue
            if column.kind in _NUMBER_FORMATS:
                cell.number_format = _NUMBER_FORMATS[column.kind]
            if isinstance(value, str):
                cell.data_type = "s"  # never let the library interpret text as a formula
            widths[index] = max(widths[index], len(str(value)))

    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = min(max(width + 2, 8), 50)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def render(
    export_format: ExportFormat,
    *,
    name: str,
    columns: Sequence[Column],
    rows: Iterable[Mapping[str, Any]],
    on_date: date,
) -> ExportFile:
    """Build the downloadable file. `name` is a slug like "products"; `on_date` is the shop's local date."""
    stem = f"{name}_{on_date.isoformat()}"
    if export_format is ExportFormat.CSV:
        return ExportFile(f"{stem}.csv", CSV_MEDIA_TYPE, render_csv(columns, rows))
    return ExportFile(
        f"{stem}.xlsx", XLSX_MEDIA_TYPE, render_xlsx(columns, rows, sheet_name=name.replace("_", " ").title())
    )


# --- Report documents (title, shop, period, generated time, notes) --------------------------------------------------


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def render_report_csv(
    title: str,
    meta: Sequence[tuple[str, str]],
    columns: Sequence[Column],
    rows: Iterable[Mapping[str, Any]],
    notes: Sequence[str],
) -> bytes:
    """CSV with a small header block (title, shop, period, generated time), a blank line, the table, then the notes."""
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([neutralize_formula(title)])
    for label, value in meta:
        writer.writerow([label, neutralize_formula(value)])
    writer.writerow([])
    writer.writerow([neutralize_formula(c.header) for c in columns])
    for row in rows:
        writer.writerow([_csv_text(row.get(c.key), c.kind) for c in columns])
    if notes:
        writer.writerow([])
        for note in notes:
            writer.writerow([neutralize_formula(note)])
    return (CSV_BOM + buffer.getvalue()).encode("utf-8")


def render_report_xlsx(
    title: str, meta: Sequence[tuple[str, str]], columns: Sequence[Column], rows: Iterable[Mapping[str, Any]], notes: Sequence[str],
    *, sheet_name: str,
) -> bytes:  # fmt: skip
    """XLSX: numbers and dates are real cells, the header row is frozen, widths fit the content, notes sit under the table."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name[:31]
    sheet.append([neutralize_formula(title)])
    sheet["A1"].font = Font(bold=True, size=14)
    for label, value in meta:
        sheet.append([label, neutralize_formula(value)])
        sheet.cell(row=sheet.max_row, column=1).font = Font(bold=True)
    sheet.append([])
    sheet.append([neutralize_formula(c.header) for c in columns])
    header_row = sheet.max_row  # the blank row above is not counted by max_row until something is written after it
    for cell in sheet[header_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", start_color="E2F0E8")
    sheet.freeze_panes = f"A{header_row + 1}"
    widths = [len(c.header) for c in columns]
    for row in rows:
        values = [_clean(row.get(c.key)) for c in columns]
        sheet.append(values)
        for i, (c, value) in enumerate(zip(columns, values, strict=True)):
            cell = sheet.cell(row=sheet.max_row, column=i + 1)
            if value is None:
                continue
            if c.kind in _NUMBER_FORMATS:
                cell.number_format = _NUMBER_FORMATS[c.kind]
            if isinstance(value, str):
                cell.data_type = "s"
            widths[i] = max(widths[i], len(str(value)))
    if notes:
        sheet.append([])
        for note in notes:
            sheet.append([neutralize_formula(note)])
            sheet.cell(row=sheet.max_row, column=1).font = Font(italic=True)
    for i, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 50)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# A small, dependency-free PDF writer: A4 landscape, a monospaced font (so columns line up and text can be cut to fit exactly),
# repeated header row, page numbers. Only Latin-1 text is drawable with the standard fonts; other scripts show as '?'.
_PAGE_W, _PAGE_H, _MARGIN = 842, 595, 36
_FONT, _CHAR_W, _LINE = 8, 4.8, 11.5


def _pdf_escape(text: str) -> str:
    safe = text.encode("latin-1", "replace").decode("latin-1")
    return safe.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def _wrap(text: str, width: int) -> list[str]:
    words, lines, line = text.split(), [], ""
    for word in words:
        while len(word) > width:
            if line:
                lines.append(line)
                line = ""
            lines.append(word[:width])
            word = word[width:]
        if len(line) + len(word) + (1 if line else 0) > width:
            lines.append(line)
            line = word
        else:
            line = f"{line} {word}".strip()
    if line:
        lines.append(line)
    return lines


_PLACES = {Kind.MONEY: "0.01", Kind.PERCENT: "0.01", Kind.QUANTITY: "0.001"}


def _pdf_text(value: Any, kind: Kind) -> str:
    if isinstance(value, Decimal) and kind in _PLACES:
        return format(value.quantize(Decimal(_PLACES[kind])), "f")
    return _text(value)


def _put(ops: list[str]):  # noqa: ANN202
    def text(s: str, bold: bool = False, size: int = _FONT, y_pos: float = 0.0) -> None:
        ops.append(f"/F{2 if bold else 1} {size} Tf 1 0 0 1 {_MARGIN} {y_pos:.1f} Tm ({_pdf_escape(s)}) Tj")

    return text


def render_pdf(
    title: str, meta: Sequence[tuple[str, str]], columns: Sequence[Column], rows: Iterable[Mapping[str, Any]], notes: Sequence[str],
    generated: str,
) -> bytes:  # fmt: skip
    usable = _PAGE_W - 2 * _MARGIN
    chars = int(usable / _CHAR_W)
    data = [[_pdf_text(r.get(c.key), c.kind) for c in columns] for r in rows]
    numeric = [c.kind not in (Kind.TEXT, Kind.DATE, Kind.DATETIME) for c in columns]
    want = [
        min(max(len(c.header), *(len(r[i]) for r in data), 4) if data else len(c.header), 40)
        for i, c in enumerate(columns)
    ]
    gaps = max(len(columns) - 1, 0) * 2
    while (
        sum(want) + gaps > chars and max(want, default=0) > 6
    ):  # shrink the widest column until the row fits the page
        want[want.index(max(want))] -= 1
    widths = want

    def cell(text: str, i: int) -> str:
        text = text if len(text) <= widths[i] else text[: max(widths[i] - 2, 1)] + ".."
        return text.rjust(widths[i]) if numeric[i] else text.ljust(widths[i])

    def line(values: list[str]) -> str:
        return "  ".join(cell(v, i) for i, v in enumerate(values))

    header = line([c.header for c in columns])
    head_lines: list[tuple[str, bool]] = [
        (title, True),
        *[(f"{k}: {v}", False) for k, v in meta],
        (generated, False),
        ("", False),
    ]
    body_top = _PAGE_H - _MARGIN
    per_page_first = int((body_top - _MARGIN - 14 - _LINE * (len(head_lines) + 2)) / _LINE)
    per_page = int((body_top - _MARGIN - 14 - _LINE * 2) / _LINE)
    note_lines = [n for note in notes for n in (_wrap(note, chars) or [""])]
    pages: list[list[str]] = []
    remaining = list(data)
    while True:
        take = per_page_first if not pages else per_page
        chunk, remaining = remaining[:take], remaining[take:]
        pages.append([line(r) for r in chunk])
        if not remaining:
            break
    if note_lines:
        capacity = per_page_first if len(pages) == 1 else per_page
        if len(pages[-1]) + len(note_lines) + 1 > capacity:
            pages.append([])  # the notes would not fit under the last row: give them their own page
    total = len(pages)
    objects: list[bytes] = []

    def add(body: bytes) -> int:
        objects.append(body)
        return len(objects)

    font_regular = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Courier /Encoding /WinAnsiEncoding >>")
    font_bold = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Courier-Bold /Encoding /WinAnsiEncoding >>")
    pages_id = len(objects) + 1 + 2 * total  # the /Pages object is added after the page and content objects
    page_ids: list[int] = []
    for n, page_rows in enumerate(pages):
        ops = ["BT"]
        y = body_top
        text = _put(ops)

        if n == 0:
            for s, bold in head_lines:
                y -= _LINE + (3 if bold else 0)
                text(s, bold, 12 if bold else _FONT, y)
        else:
            y -= _LINE
            text(f"{title} (continued)", True, 10, y)
            y -= _LINE
        y -= _LINE
        text(header, True, _FONT, y)
        for r in page_rows:
            y -= _LINE
            text(r, False, _FONT, y)
        if n == total - 1 and note_lines:
            y -= _LINE
            for s in note_lines:
                y -= _LINE
                text(s, False, _FONT, y)
        text(f"Page {n + 1} of {total}", False, 7, _MARGIN - 14 + 6)
        ops.append("ET")
        stream = "\n".join(ops).encode("latin-1", "replace")
        content = add(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
        page_ids.append(
            add(
                f"<< /Type /Page /Parent {pages_id} 0 R /MediaBox [0 0 {_PAGE_W} {_PAGE_H}] /Contents {content} 0 R "
                f"/Resources << /Font << /F1 {font_regular} 0 R /F2 {font_bold} 0 R >> >> >>".encode()
            )
        )
    kids = " ".join(f"{i} 0 R" for i in page_ids)
    assert add(f"<< /Type /Pages /Kids [{kids}] /Count {total} >>".encode()) == pages_id
    catalog = add(f"<< /Type /Catalog /Pages {pages_id} 0 R >>".encode())
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode() + b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root {catalog} 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)
